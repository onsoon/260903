"""BeAT 1차시 실습 자료 생성기: beat-class/data/ 에 docx, txt, csv 를 만든다."""
import csv, random, sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent.parent / "data"
OUT.mkdir(parents=True, exist_ok=True)

# ---------- 1. 수업GPT: 통합과학 학습지 ----------
SCI_TITLE = "[통합과학] 탄소 순환과 기후 변화 학습지"
SCI = [
    ("h", "1. 탄소는 어디에 있을까?"),
    ("p", "지구의 탄소는 네 곳에 나뉘어 있다. 이 네 곳을 ‘탄소 저장소’라고 부른다."),
    ("ul", ["대기: 이산화 탄소(CO₂), 메테인(CH₄) 같은 기체",
            "생물권: 생물의 몸을 이루는 유기물(탄수화물, 단백질, 지방)",
            "수권(해양): 물에 녹아 있는 이산화 탄소, 탄산 이온, 조개껍데기의 탄산 칼슘",
            "지권: 석회암 같은 탄산염 암석, 석탄·석유·천연가스 같은 화석 연료"]),
    ("h", "2. 탄소는 어떻게 이동할까?"),
    ("ul", ["광합성: 식물이 대기의 이산화 탄소를 흡수해 유기물을 만든다. (대기 → 생물권)",
            "호흡: 생물이 유기물을 분해해 에너지를 얻고 이산화 탄소를 내보낸다. (생물권 → 대기)",
            "분해: 미생물이 사체와 배설물을 분해해 이산화 탄소를 내보낸다. (생물권 → 대기, 토양)",
            "연소: 화석 연료를 태우면 오랫동안 땅속에 있던 탄소가 이산화 탄소가 되어 대기로 나온다. (지권 → 대기)",
            "해양의 흡수와 방출: 바다는 대기의 이산화 탄소를 녹여 흡수하기도 하고 내보내기도 한다. (대기 ↔ 수권)",
            "퇴적: 생물의 유해와 탄산염이 아주 오랜 시간 쌓여 암석과 화석 연료가 된다. (생물권·수권 → 지권)"]),
    ("h", "3. 온실 효과"),
    ("p", "태양 빛을 받아 데워진 지표는 적외선을 내보낸다. 대기 중의 온실 기체(수증기, 이산화 탄소, 메테인 등)는 이 적외선의 일부를 흡수했다가 다시 지표 쪽으로 내보낸다. 이 덕분에 지구의 평균 기온은 약 15 °C로 유지된다. 온실 효과가 없다면 지구 평균 기온은 약 -18 °C까지 내려갈 것으로 계산된다."),
    ("p", "즉, 온실 효과 자체는 생명이 살 수 있게 해 주는 자연스러운 현상이다."),
    ("h", "4. 지구 온난화"),
    ("p", "산업 혁명 이후 화석 연료 사용과 숲의 감소로 대기 중 온실 기체가 늘어났다. 대기 중 이산화 탄소 농도는 산업 혁명 이전 약 280 ppm에서 최근 420 ppm을 넘었다. 온실 기체가 늘어 온실 효과가 강해지면서 지구 평균 기온이 올라가는 현상을 지구 온난화라고 한다."),
    ("ul", ["해수면 상승: 빙하가 녹고, 바닷물이 데워지며 부피가 늘어난다(열팽창).",
            "이상 기후: 폭염, 집중 호우, 가뭄이 더 자주, 더 강하게 나타난다.",
            "생태계 변화: 생물이 사는 지역이 바뀌고, 개화 시기가 달라진다.",
            "해양 산성화: 바다에 녹는 이산화 탄소가 늘어 바닷물이 산성 쪽으로 변한다."]),
    ("h", "5. 자주 헷갈리는 개념"),
    ("ul", ["온실 효과 ≠ 지구 온난화: 온실 효과는 원래 있던 자연 현상이고, 지구 온난화는 온실 효과가 강해져 기온이 올라가는 현상이다.",
            "광합성과 호흡은 정반대 방향의 이동이다. 식물도 호흡을 한다.",
            "화석 연료의 탄소는 원래 생물의 몸에 있던 탄소다."]),
    ("h", "6. 확인 문제"),
    ("ol", ["대기의 탄소를 생물권으로 옮기는 과정은 무엇인가?",
            "‘온실 효과는 나쁜 현상이므로 없어져야 한다.’ 이 문장의 잘못된 점을 설명하시오.",
            "해수면이 상승하는 원인 두 가지를 쓰시오."]),
    ("p", "정답: 1. 광합성  2. 온실 효과가 없으면 지구 평균 기온이 약 -18 °C가 되어 생물이 살기 어렵다. 문제는 온실 효과가 지나치게 강해진 것(지구 온난화)이다.  3. 빙하가 녹아 바다로 흘러드는 것, 바닷물이 데워져 부피가 늘어나는 것(열팽창)"),
]

# ---------- 2. 개인GPT: 수학 오답노트 ----------
MATH_TITLE = "[공통수학] 나의 오답노트 (중간고사 대비)"
MATH = [
    ("p", "※ 아래 풀이에는 일부러 틀린 단계가 들어 있어요. AI에게 정답을 바로 묻지 말고, 어느 단계가 틀렸는지 힌트를 받아 스스로 고쳐 보세요."),
    ("h", "오답 1. 이차방정식 (근의 공식)"),
    ("p", "문제: 이차방정식 x² - 6x + 5 = 0 의 두 근을 구하시오."),
    ("ol", ["a = 1, b = -6, c = 5",
            "근의 공식 x = (-b ± √(b² - 4ac)) / 2a 에 대입한다.",
            "x = (-6 ± √(36 - 20)) / 2",
            "x = (-6 ± 4) / 2",
            "내 답: x = -1 또는 x = -5"]),
    ("h", "오답 2. 나머지정리"),
    ("p", "문제: 다항식 P(x) = x³ - 2x² + 3x - 4 를 x - 2 로 나눈 나머지를 구하시오."),
    ("ol", ["나머지정리에 따라 나머지는 P(-2) 이다.",
            "P(-2) = (-2)³ - 2·(-2)² + 3·(-2) - 4",
            "= -8 - 8 - 6 - 4",
            "내 답: 나머지는 -26"]),
    ("h", "오답 3. 이차부등식"),
    ("p", "문제: 이차부등식 x² - x - 6 < 0 의 해를 구하시오."),
    ("ol", ["좌변을 인수분해하면 (x - 3)(x + 2) < 0",
            "x = 3, x = -2 가 경계이다.",
            "내 답: x < -2 또는 x > 3"]),
    ("h", "정답 확인용 (AI와 먼저 풀어 본 뒤에 보기)"),
    ("ul", ["오답 1: x = 1 또는 x = 5 (3단계에서 -b 는 +6 이어야 함)",
            "오답 2: 나머지는 P(2) = 2 (x - a 로 나눈 나머지는 P(a))",
            "오답 3: -2 < x < 3 (곱이 음수이면 두 근 사이)"]),
]

def make_docx(path, title, blocks, foot):
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Malgun Gothic"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    st.font.size = Pt(11)
    h = d.add_heading(title, level=1)
    for r in h.runs:
        r.font.color.rgb = RGBColor(0x21, 0x60, 0xD6)
        r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "맑은 고딕")
    d.add_paragraph(foot).runs[0].font.size = Pt(9)
    for kind, val in blocks:
        if kind == "h":
            p = d.add_heading(val, level=2)
            for r in p.runs:
                r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "맑은 고딕")
        elif kind == "p":
            d.add_paragraph(val)
        elif kind in ("ul", "ol"):
            for item in val:
                d.add_paragraph(item, style="List Bullet" if kind == "ul" else "List Number")
    d.save(path)

def make_txt(path, title, blocks, foot):
    lines = [title, foot, ""]
    for kind, val in blocks:
        if kind == "h":
            lines += ["", val]
        elif kind == "p":
            lines.append(val)
        elif kind == "ul":
            lines += [f"- {v}" for v in val]
        elif kind == "ol":
            lines += [f"{i}. {v}" for i, v in enumerate(val, 1)]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")

FOOT = "BeAT 1차시 실습용 예시 자료 · 수업 연습을 위해 만든 자료입니다."
make_docx(OUT / "수업GPT_통합과학_탄소순환_학습지.docx", SCI_TITLE, SCI, FOOT)
make_txt(OUT / "수업GPT_통합과학_탄소순환_학습지.txt", SCI_TITLE, SCI, FOOT)
make_docx(OUT / "개인GPT_수학_오답노트.docx", MATH_TITLE, MATH, FOOT)
make_txt(OUT / "개인GPT_수학_오답노트.txt", MATH_TITLE, MATH, FOOT)

# ---------- 3. 개인GPT: 비문학 지문 ----------
(OUT / "개인GPT_비문학지문_추천알고리즘.txt").write_text("""[비문학 지문] 내 피드는 왜 나만 보는 것으로 가득할까
BeAT 1차시 실습용 예시 자료 · 수업 연습을 위해 새로 쓴 지문입니다.

동영상 앱을 켜면 첫 화면에 내가 좋아할 만한 영상이 줄지어 나온다. 어제 고양이 영상을 두 개 봤다면 오늘은 고양이 영상이 열 개쯤 보인다. 이것은 우연이 아니라 추천 알고리즘이 작동한 결과다.

추천 알고리즘은 사용자가 남긴 흔적을 모아 다음에 보여 줄 콘텐츠를 고른다. 흔적에는 무엇을 클릭했는지, 영상을 몇 초까지 봤는지, ‘좋아요’를 눌렀는지, 어떤 시간대에 앱을 켰는지 같은 정보가 들어간다. 알고리즘은 이런 기록이 비슷한 다른 사용자들이 무엇을 좋아했는지도 함께 살핀다. 나와 취향이 비슷한 사람들이 즐겨 본 영상이라면 나도 좋아할 가능성이 높다고 판단하는 것이다. 이렇게 하면 사용자는 원하는 콘텐츠를 빨리 찾을 수 있고, 서비스는 사용자가 더 오래 머물게 할 수 있다.

그러나 편리함에는 그림자가 있다. 알고리즘이 내가 이미 좋아하는 것만 계속 보여 주면, 나는 점점 비슷한 정보와 의견에만 둘러싸이게 된다. 이처럼 맞춤형 정보가 보이지 않는 막처럼 사용자를 감싸 다른 정보를 걸러 내는 현상을 ‘필터 버블’이라고 부른다. 필터 버블 안에서는 내 생각과 다른 의견을 만날 기회가 줄어들고, 세상 사람들 대부분이 나와 같은 생각을 한다고 착각하기 쉽다. 한쪽으로 치우친 정보만 반복해서 보면 생각이 더 굳어지기도 한다.

필터 버블에서 벗어나려면 무엇이 필요할까? 먼저 내가 보는 화면이 ‘모두에게 같은 화면’이 아니라는 사실을 아는 것이 출발점이다. 그다음 일부러 다른 분야의 콘텐츠를 찾아보거나, 같은 사건을 다룬 여러 매체의 기사를 비교해 보는 습관이 도움이 된다. 시청 기록을 정리하거나 추천 설정을 바꾸는 것도 방법이다. 알고리즘은 내가 남긴 흔적으로 움직이므로, 내가 어떤 흔적을 남길지 스스로 고르는 것이 곧 내 피드를 바꾸는 일이다.

추천 알고리즘은 그 자체로 좋거나 나쁜 도구가 아니다. 도구를 쓰는 사람이 그 원리를 알고 있을 때, 편리함은 누리면서 그림자는 줄일 수 있다.
""", encoding="utf-8")

# ---------- 4. AI Agents: 가상 설문 데이터 ----------
rnd = random.Random(20261018)
rows = []
for i in range(1, 41):
    grade = rnd.choice([1, 1, 2, 2, 3])
    phone = round(min(8.5, max(1.0, rnd.gauss(4.2, 1.6))), 1)
    bed = int(min(120, max(0, rnd.gauss(phone * 11, 15))) // 5 * 5)
    sleep = round(min(8.5, max(4.5, 7.6 - 0.018 * bed - 0.12 * (grade - 1) + rnd.gauss(0, 0.45))), 1)
    tired = int(min(5, max(1, round(5.2 - (sleep - 4.5) * 0.85 + rnd.gauss(0, 0.6)))))
    rows.append([i, grade, phone, bed, sleep, tired])
with open(OUT / "AIAgents_가상설문_스마트폰과수면.csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f)
    w.writerow(["번호", "학년", "하루_스마트폰_사용시간(시간)", "잠들기전_스마트폰_사용(분)", "평일_수면시간(시간)", "아침_피곤함(1~5)"])
    w.writerows(rows)

# ---------- 5. Custom AI: 자기소개 초안 ----------
(OUT / "CustomAI_자기소개_초안_예시.txt").write_text("""[자기소개 초안 예시] 학생 A (가상 인물)
BeAT 1차시 실습용 예시 자료 · 실제 학생의 글이 아닙니다. 피드백 연습용으로 일부러 아쉬운 점을 넣었습니다.

저는 어릴 때부터 성실하고 책임감이 강한 학생입니다. 고등학교에 들어와서 다양한 활동을 하며 많은 것을 배웠습니다. 특히 과학 동아리에서 친구들과 함께 활동하면서 협동심의 중요성을 느꼈습니다. 동아리에서는 학교 주변 미세먼지를 측정하는 간이 측정기를 만들었는데, 처음에는 값이 너무 들쭉날쭉해서 센서 위치를 세 번이나 바꿔야 했습니다. 이 경험을 통해 저는 포기하지 않는 끈기를 배웠습니다. 저는 앞으로도 열심히 노력하여 꼭 꿈을 이루는 사람이 되겠습니다.

[피드백을 받을 때 생각해 볼 점]
- 어떤 문장이 누구나 쓸 수 있는 말일까?
- 가장 구체적이고 나다운 경험은 어느 문장에 있을까?
- 그 경험을 앞으로 끌어오면 글이 어떻게 달라질까?
""", encoding="utf-8")

# ---------- 6. Custom AI: 영어 지문 ----------
(OUT / "CustomAI_영어지문_번역연습.txt").write_text("""[English Reading] Why Are Teenagers So Sleepy in the Morning?
BeAT 1차시 실습용 예시 자료 · 수업 연습을 위해 새로 쓴 영어 지문입니다.

Many teenagers find it hard to fall asleep before midnight and even harder to get up for school. This is not simply laziness. During adolescence, the body's internal clock, often called the circadian rhythm, shifts later. The brain starts releasing melatonin, a hormone that makes us feel sleepy, later in the evening than it did in childhood. As a result, a teenager may not feel tired until late at night.

The problem is that school still starts early. When students go to bed late and wake up early, they lose hours of sleep every week. Lack of sleep can make it harder to concentrate, remember new information, and control emotions.

Experts suggest a few simple habits: keep a regular sleep schedule, even on weekends, and put screens away before bed, because bright light in the evening can delay melatonin even more.

[번역 연습 과제]
1. Custom AI의 번역 도구로 위 지문을 번역해 보세요.
2. 원문과 번역문을 한 문장씩 비교해서, 어색하거나 뜻이 달라진 문장을 하나 찾아보세요.
3. 지문의 핵심 단어 3개(예: circadian rhythm)를 골라 뜻을 정리하세요.
""", encoding="utf-8")

print("\n".join(sorted(p.name for p in OUT.iterdir())))
